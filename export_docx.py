"""Export the assembled resume (and cover letter) from a run directory to .docx
in the layout of the author's original resume (resume/base_resume.docx).

    python export_docx.py --run outputs/<run>            # resume_tailored*.md -> <Name>_Resume_<jd>_<time>.docx
    python export_docx.py --run outputs/<run> --letter   # also cover letter from cv_suggestions*.json

Layout reproduced from the original: A4 with a full-width navy header band
(name 22pt white, title line in cyan, three contact columns in pale blue),
navy section banners with cyan "▐  TITLE" text, a three-column education
table (dates | university | degree), project title rows with bullet items
whose leading "Label:" is bold navy, "◆ Category  items" skill lines, and a
two-column block for additional projects. Only the resume body is exported;
the source appendix after the '---' rule stays in the Markdown.
"""
import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from job_identity import load_identity
from tailor_cv import latest_output, output_number

PROJECT_ROOT = Path(__file__).resolve().parent
FONT = "Calibri"
NAVY = RGBColor(0x1A, 0x3A, 0x5C)
CYAN = RGBColor(0x00, 0xA8, 0xCC)
BLUE = RGBColor(0x1F, 0x5F, 0xA6)
PALE = RGBColor(0xA9, 0xC8, 0xE0)
GREY = RGBColor(0x88, 0x88, 0x88)
INK = RGBColor(0x1C, 0x1C, 0x1C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_HEX = "1A3A5C"
PAGE_WIDTH_CM = 21.0
MARGIN_CM = 1.1
CONTENT_CM = PAGE_WIDTH_CM - 2 * MARGIN_CM
CONTACT_PREFIXES = ("📍", "✉", "📞", "⌨", "🔗", "☎", "✆")
LABEL = re.compile(r"^([A-Z][A-Za-z0-9 &/+\-]{1,28}):\s+(.*)$")


# ---------------------------------------------------------------- markdown


def resume_body(markdown):
    """Lines before the '---' separator that precedes the source appendix."""
    body = markdown.split("\n---\n", 1)[0]
    return [line.rstrip() for line in body.splitlines()]


def parse_resume(lines):
    """-> {name, title, contact[], sections: [{title, blocks: [{heading, lines[]}]}]}"""
    doc = {"name": "", "title": "", "contact": [], "sections": []}
    section = None
    block = None
    in_header = True
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# ") and not doc["name"]:
            doc["name"] = line[2:].strip()
            continue
        if line.startswith("## "):
            in_header = False
            section = {"title": line[3:].strip(), "blocks": []}
            block = None
            doc["sections"].append(section)
            continue
        if in_header:
            if line.startswith(CONTACT_PREFIXES):
                doc["contact"].append(line)
            elif not doc["title"]:
                doc["title"] = line
            else:
                doc["contact"].append(line)
            continue
        if section is None:
            continue
        if line.startswith("### "):
            block = {"heading": line[4:].strip(), "lines": []}
            section["blocks"].append(block)
            continue
        if block is None:
            block = {"heading": None, "lines": []}
            section["blocks"].append(block)
        kind = "bullet" if re.match(r"^[-*]\s+", line) else "text"
        block["lines"].append({"kind": kind, "text": re.sub(r"^[-*]\s+", "", line)})
    return doc


# ---------------------------------------------------------------- docx helpers


def _shade(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    # schema order: tcW, gridSpan, vMerge, tcBorders, shd, noWrap, tcMar, textDirection, tcFitText, vAlign, hideMark
    tc_pr.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
                                "w:vAlign", "w:hideMark")


def _cell_margins(cell, top=0, bottom=0, left=0, right=0):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(int(value * 20)))  # points -> twips
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.insert_element_before(mar, "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")


def _no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    # schema order: ... tblInd, tblBorders, shd, tblLayout, tblCellMar, tblLook ...
    tbl_pr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                                 "w:tblCaption", "w:tblDescription")


def _table(doc, widths_cm, fill=None):
    table = doc.add_table(rows=1, cols=len(widths_cm))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _no_borders(table)
    for cell, width in zip(table.rows[0].cells, widths_cm):
        cell.width = Cm(width)
        if fill:
            _shade(cell, fill)
    return table


def _run(paragraph, text, size=8, bold=False, color=INK, italic=False, spacing=None):
    run = paragraph.add_run(text)
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    if spacing:  # character spacing in twentieths of a point, as in the original banners
        node = OxmlElement("w:spacing")
        node.set(qn("w:val"), str(spacing))
        # schema order: ... color, spacing, w, kern, position, sz, ...
        run._element.rPr.insert_element_before(
            node, "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u",
            "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs",
            "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")
    return run


def _para(container, before=0, after=0, align=None, style=None):
    para = container.add_paragraph(style=style) if style else container.add_paragraph()
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.05
    if align is not None:
        para.alignment = align
    return para


def _first_para(cell):
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    return para


def _bullet(container, text, label_bold=True):
    """Bullet item; a leading 'Label:' is rendered bold navy like the original."""
    para = _para(container, before=0.45, after=0.45, style="List Bullet")
    para.paragraph_format.left_indent = Cm(0.75)
    para.paragraph_format.first_line_indent = Cm(-0.35)
    match = LABEL.match(text) if label_bold else None
    if match:
        _run(para, match.group(1) + ": ", bold=True, color=NAVY)
        _run(para, match.group(2))
    else:
        _run(para, text)
    return para


def _banner(doc, title):
    gap = doc.add_paragraph()  # breathing room above the banner, as in the original
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.space_after = Pt(0)
    gap.paragraph_format.line_spacing = 0.5
    gap.paragraph_format.keep_with_next = True
    table = _table(doc, [CONTENT_CM], fill=NAVY_HEX)
    cell = table.rows[0].cells[0]
    _cell_margins(cell, top=2.2, bottom=2.2, left=4, right=4)
    para = _first_para(cell)
    para.paragraph_format.keep_with_next = True
    _run(para, f"▐  {title.upper()}", size=9, bold=True, color=CYAN, spacing=60)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.line_spacing = 0.6
    spacer.paragraph_format.keep_with_next = True


def _header(doc, parsed):
    table = _table(doc, [CONTENT_CM + 2 * MARGIN_CM], fill=NAVY_HEX)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    _cell_margins(cell, top=9, bottom=8, left=20, right=20)
    name = _first_para(cell)
    name.paragraph_format.space_after = Pt(1)
    _run(name, parsed["name"], size=22, bold=True, color=WHITE)
    if parsed["title"]:
        title = _para(cell, after=2.5)
        _run(title, parsed["title"], size=7.5, color=CYAN)
    contact = parsed["contact"]
    if contact:
        columns = [contact[0:1], contact[1:2], contact[2:]]
        inner = cell.add_table(rows=1, cols=3)
        _no_borders(inner)
        inner.autofit = False
        for col, (icell, width) in zip(columns, zip(inner.rows[0].cells, (5.2, 6.4, 7.0))):
            icell.width = Cm(width)
            _shade(icell, NAVY_HEX)
            first = True
            for line in col:
                para = _first_para(icell) if first else _para(icell)
                first = False
                _run(para, line, size=7.5, color=PALE)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 0.6


def _split_degree_line(text):
    """Separate dates in pipe or trailing-parenthesis format; preserve wording."""
    if "|" in text:
        left, right = [part.strip() for part in text.rsplit("|", 1)]
        if re.search(r"\d{4}", right):
            return left, right
        if re.search(r"\d{4}", left):
            return right, left
    match = re.fullmatch(r"(.+?)\s*[（(]([^()（）]*\d{4}[^()（）]*)[）)]\s*", text)
    if match:
        return match.group(1).strip(), match.group(2).strip()
    return text, ""


def _education(doc, section):
    widths = (4.2, 7.8, 6.8)
    table = _table(doc, widths)
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    table.rows[0]._tr.getparent().remove(table.rows[0]._tr)
    for block in section["blocks"]:
        school = block["heading"] or ""
        degree, dates, extra = "", "", []
        for line in block["lines"]:
            if school and not degree and line["kind"] == "text" and not LABEL.match(line["text"]):
                degree, dates = _split_degree_line(line["text"])
            else:
                extra.append(line["text"])
        row = table.add_row()
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        p0 = _first_para(row.cells[0]); p0.paragraph_format.space_before = Pt(2)
        _run(p0, dates, size=8.5, bold=True, color=NAVY)
        p1 = _first_para(row.cells[1], ); p1.paragraph_format.space_before = Pt(2)
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p1, school, size=8.5, bold=True, color=NAVY)
        p2 = _first_para(row.cells[2]); p2.paragraph_format.space_before = Pt(2)
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _run(p2, degree, size=8.5, bold=True, color=BLUE)
        for text in extra:
            # Keep notes next to their own school, not after the entire table.
            note_row = table.add_row()
            cell = note_row.cells[0].merge(note_row.cells[-1])
            para = _first_para(cell)
            para.paragraph_format.space_before = Pt(0.3)
            para.paragraph_format.space_after = Pt(1)
            match = LABEL.match(text)
            if match:
                _run(para, match.group(1) + ": ", size=8, bold=True, color=NAVY)
                _run(para, match.group(2), size=8, bold=True, color=NAVY)
            else:
                _run(para, text, size=8, bold=True, color=NAVY)


def _project_block(container, block, title_size=9, keep_heading=True):
    if block["heading"]:
        para = _para(container, before=2.4, after=0.3)
        # Inside a table Word reads keep-with-next as "keep this row with the
        # next row", which glues the whole table together; only use it in body text.
        para.paragraph_format.keep_with_next = keep_heading
        _run(para, block["heading"], size=title_size, bold=True, color=NAVY)
    for line in block["lines"]:
        if line["kind"] == "bullet":
            _bullet(container, line["text"])
        else:
            para = _para(container, before=0, after=0.5)
            _run(para, line["text"], size=8, color=BLUE)


def _two_columns(doc, blocks, renderer):
    """Two blocks per table row. One tall row holding everything would not
    break across pages in Word and jumps to the next page as a whole, leaving
    a large gap; short rows flow with the text."""
    half = CONTENT_CM / 2
    table = _table(doc, [half, half])
    table.rows[0]._tr.getparent().remove(table.rows[0]._tr)
    for start in range(0, len(blocks), 2):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))  # no orphaned last line
        pair = blocks[start:start + 2]
        if len(pair) == 1:  # an odd last block takes the full width: fewer lines than half a column
            cells = [row.cells[0].merge(row.cells[1])]
            cells[0].width = Cm(CONTENT_CM)
        else:
            cells = list(row.cells)
            for cell in cells:
                cell.width = Cm(half)
        for cell, block in zip(cells, pair):
            _cell_margins(cell, left=2, right=6)
            for stray in list(cell.paragraphs):
                stray._p.getparent().remove(stray._p)
            renderer(cell, block)
    return table


def _skills(doc, section):
    for block in section["blocks"]:
        for line in block["lines"]:
            text = line["text"]
            match = LABEL.match(text)
            para = _para(doc, before=0.65, after=0.65)
            if match:
                _run(para, "◆ " + match.group(1) + "  ", size=8, bold=True, color=NAVY)
                _run(para, match.group(2), size=8, bold=True, color=NAVY)
            else:
                _run(para, ("" if text.startswith("◆") else "◆ ") + text, size=8, bold=True, color=NAVY)


def _generic(doc, section, bullets_bold_label=True):
    for block in section["blocks"]:
        if block["heading"]:
            para = _para(doc, before=1.5, after=0.4)
            _run(para, block["heading"], size=8.5, bold=True, color=NAVY)
        for line in block["lines"]:
            if line["kind"] == "bullet":
                _bullet(doc, line["text"], bullets_bold_label)
            else:
                para = _para(doc, before=0.3, after=1.2)
                match = LABEL.match(line["text"])
                if match and bullets_bold_label:
                    _run(para, match.group(1) + ": ", size=8, bold=True, color=NAVY)
                    _run(para, match.group(2), size=8)
                else:
                    _run(para, line["text"], size=8)


def render_resume(parsed, out_path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(PAGE_WIDTH_CM), Cm(29.7)
    section.left_margin = section.right_margin = Cm(MARGIN_CM)
    section.top_margin, section.bottom_margin = Cm(0.9), Cm(0.9)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(8)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for stray in doc.paragraphs:  # python-docx may start with an empty paragraph
        stray._p.getparent().remove(stray._p)

    _header(doc, parsed)
    for sec in parsed["sections"]:
        key = sec["title"].strip().upper()
        _banner(doc, sec["title"])
        if key == "EDUCATION":
            _education(doc, sec)
        elif "PROJECT" in key and ("ADDITIONAL" in key or "OTHER" in key):
            _two_columns(doc, sec["blocks"],
                         lambda cell, block: _project_block(cell, block, title_size=8.5,
                                                            keep_heading=False))
        elif "PROJECT" in key or "EXPERIENCE" in key:
            for block in sec["blocks"]:
                _project_block(doc, block)
        elif "SKILL" in key and "LANGUAGE" not in key:
            _skills(doc, sec)
        else:
            _generic(doc, sec)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def markdown_to_docx(lines, out_path):
    return render_resume(parse_resume(lines), out_path)


def cover_letter_to_docx(suggestions, out_path, name_line=None):
    paragraphs = [p["text"] for p in suggestions.get("cover_letter", []) if p.get("text")]
    if not paragraphs:
        return None
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(PAGE_WIDTH_CM), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(10.5)
    for stray in doc.paragraphs:  # python-docx may start with an empty paragraph
        stray._p.getparent().remove(stray._p)
    if name_line:
        head = _para(doc, after=10)
        _run(head, name_line, size=16, bold=True, color=NAVY)
    for text in paragraphs:
        for chunk in text.split("\n"):
            if chunk.strip():
                para = _para(doc, after=7)
                para.paragraph_format.line_spacing = 1.2
                _run(para, chunk.strip(), size=10.5)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


DOCX_KINDS = {"resume": "Resume", "cover_letter": "Cover_Letter"}
LEGACY_DOCX = {"resume": "resume_tailored", "cover_letter": "cover_letter"}


def _file_part(text, limit):
    """English letters, digits and hyphens only: such a name survives email
    attachments, upload forms and applicant tracking systems unchanged."""
    cleaned = re.sub(r"[^A-Za-z0-9]+", "-", text or "")
    return cleaned.strip("-")[:limit].rstrip("-")


_LEGAL_SUFFIX = re.compile(r"[\s,]+(?:Pty\.?\s+Ltd\.?|Limited|Ltd\.?|Inc\.?|LLC|Corporation|Corp\.?)\s*$",
                           re.IGNORECASE)


def _run_time(run_dir):
    stamp = re.match(r"\d{4}(\d{4})-(\d{4})", Path(run_dir).name)
    return f"{stamp.group(1)}-{stamp.group(2)}" if stamp else _file_part(Path(run_dir).name, 20)


def run_label(run_dir, label=None):
    """What tells one application's files from another's, company first:
    a label given by the user, else the company and job title found in the JD
    (job_identity.json), else the English part of the JD file name; always
    followed by when the run started."""
    run_dir = Path(run_dir)
    who = _file_part(label, 40)
    if not who:
        identity = load_identity(run_dir)
        company = _LEGAL_SUFFIX.sub("", identity.get("company") or "")
        who = "_".join(part for part in (_file_part(company, 30),
                                         _file_part(identity.get("job_title"), 30)) if part)
    if not who:
        try:
            meta = json.loads((run_dir / "run_meta.json").read_text(encoding="utf-8"))
            jd_stem = Path(str(meta.get("input", {}).get("jd_path", "")).replace("\\", "/")).stem
        except (OSError, ValueError):
            jd_stem = ""
        who = _file_part(re.sub(r"^\d{8}-\d{4}-", "", jd_stem), 40)
    return "_".join(part for part in (who, _run_time(run_dir)) if part)


def docx_name(kind, person, run_dir, number=1, label=None):
    """e.g. Alex_Sample_Resume_Acme-Energy_Data-Engineer_0917-1459.docx (-2 for a
    re-export of a second resume version in the same run)."""
    who = "_".join(word.capitalize() for word in _file_part(person, 40).split("-") if word)
    parts = [part for part in (who, DOCX_KINDS[kind], run_label(run_dir, label)) if part]
    return "_".join(parts) + (f"-{number}" if number > 1 else "") + ".docx"


def latest_docx(run_dir, kind):
    """Newest exported .docx of a kind, including files written under the old
    fixed names (resume_tailored.docx, cover_letter.docx)."""
    run_dir = Path(run_dir)
    files = list(run_dir.glob(f"*_{DOCX_KINDS[kind]}_*.docx")) + list(run_dir.glob(f"{LEGACY_DOCX[kind]}*.docx"))
    return max(files, key=lambda path: path.stat().st_mtime) if files else None


def export_run(run_dir, with_letter=False, resume_md=None, label=None):
    run_dir = Path(run_dir)
    md_path = Path(resume_md) if resume_md else latest_output(run_dir, "resume_tailored")
    if md_path is None:
        raise ValueError(f"{run_dir} 里没有 resume_tailored*.md；先运行 build_resume.py")
    lines = resume_body(md_path.read_text(encoding="utf-8"))
    name_line = next((l[2:].strip() for l in lines if l.startswith("# ")), None)
    number = output_number(md_path, "resume_tailored")
    outputs = {"resume": markdown_to_docx(lines, run_dir / docx_name("resume", name_line, run_dir, number, label))}
    if with_letter:
        suggestions_path = latest_output(run_dir, "cv_suggestions", ".json")
        if suggestions_path is not None:
            suggestions = json.loads(suggestions_path.read_text(encoding="utf-8"))
            letter = cover_letter_to_docx(suggestions, run_dir / docx_name("cover_letter", name_line, run_dir, label=label),
                                          name_line)
            if letter is not None:
                outputs["cover_letter"] = letter
    return outputs


def main(argv=None):
    parser = argparse.ArgumentParser(description="把组装好的简历导出为 Word（沿用原版排版）")
    parser.add_argument("--run", required=True, type=Path)
    parser.add_argument("--md", type=Path, default=None, help="指定 resume_tailored*.md，默认最新")
    parser.add_argument("--letter", action="store_true", help="同时导出 Cover Letter")
    parser.add_argument("--label", default=None,
                        help="文件名里的公司/岗位标识（英文）；默认取 job_identity.json 里的公司与职位")
    args = parser.parse_args(argv)
    try:
        outputs = export_run(args.run, args.letter, args.md, args.label)
    except (OSError, ValueError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1
    for kind, path in outputs.items():
        print(f"已写入（{kind}）：{path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
